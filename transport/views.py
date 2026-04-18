import openpyxl
from io import BytesIO
from functools import wraps
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.http import JsonResponse, HttpResponse
from django.contrib import messages
from django.db.models import Q
from .models import Driver, LineSupervisor, Student, FamilyContact
from .forms import StudentForm, DriverForm, LineSupervisorForm, EmployeeForm


# ── Decorator ────────────────────────────────────────────────────────────────

def manager_required(func):
    @wraps(func)
    @login_required
    def wrapper(request, *args, **kwargs):
        if not request.user.is_staff:
            return redirect('employee_search')
        return func(request, *args, **kwargs)
    return wrapper


# ── Auth ──────────────────────────────────────────────────────────────────────

def login_view(request):
    if request.user.is_authenticated:
        return redirect('home')
    error = None
    if request.method == 'POST':
        user = authenticate(
            request,
            username=request.POST.get('username', '').strip(),
            password=request.POST.get('password', '')
        )
        if user is None:
            error = 'اسم المستخدم أو كلمة المرور غير صحيحة'
        elif not user.is_active:
            error = 'هذا الحساب معطّل. تواصل مع المدير لتفعيله.'
        else:
            login(request, user)
            return redirect('home')
    return render(request, 'transport/login.html', {'error': error})


def logout_view(request):
    logout(request)
    return redirect('login')


@login_required
def home(request):
    if request.user.is_staff:
        return redirect('manager_page')
    return redirect('employee_search')


# ── Manager: Dashboard ────────────────────────────────────────────────────────

@manager_required
def manager_dashboard(request):
    context = {
        'total_students': Student.objects.count(),
        'total_drivers': Driver.objects.count(),
        'total_supervisors': LineSupervisor.objects.count(),
        'total_employees': User.objects.filter(is_staff=False, is_active=True).count(),
        'recent_students': Student.objects.select_related('driver', 'supervisor').order_by('-id')[:10],
    }
    return render(request, 'transport/manager_dashboard.html', context)


# ── Manager: Students ─────────────────────────────────────────────────────────

@manager_required
def manager_students(request):
    q = request.GET.get('q', '').strip()
    qs = Student.objects.select_related('driver', 'supervisor').all()
    if q:
        qs = qs.filter(Q(name__icontains=q) | Q(grade__icontains=q) | Q(area__icontains=q) | Q(landmark__icontains=q))
    return render(request, 'transport/manager_students.html', {
        'students': qs, 'q': q, 'total': qs.count()
    })


@manager_required
def manager_student_form(request, pk=None):
    student = get_object_or_404(Student, pk=pk) if pk else None
    existing = list(student.contacts.order_by('order')) if student else []

    if request.method == 'POST':
        form = StudentForm(request.POST, request.FILES, instance=student)
        if form.is_valid():
            s = form.save()
            s.contacts.all().delete()
            for i in range(1, 7):
                name = request.POST.get(f'cn{i}', '').strip()
                phone = request.POST.get(f'cp{i}', '').strip()
                if name:
                    FamilyContact.objects.create(student=s, name=name, phone=phone, order=i)
            messages.success(request, 'تم الحفظ بنجاح ✓')
            return redirect('manager_students')
    else:
        form = StudentForm(instance=student)

    # Pad contacts to 6 slots
    contacts_padded = existing + [None] * (6 - len(existing))
    return render(request, 'transport/manager_student_form.html', {
        'form': form, 'student': student, 'contacts': contacts_padded,
        'drivers': Driver.objects.all(), 'supervisors': LineSupervisor.objects.all(),
    })


@manager_required
def manager_student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if request.method == 'POST':
        name = student.name
        student.delete()
        messages.success(request, f'تم حذف {name}')
        return redirect('manager_students')
    return render(request, 'transport/confirm_delete.html', {'obj': student, 'type': 'تلميذ'})


# ── Manager: Drivers ──────────────────────────────────────────────────────────

@manager_required
def manager_drivers(request):
    if request.method == 'POST':
        form = DriverForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'تم إضافة السائق ✓')
            return redirect('manager_drivers')
    else:
        form = DriverForm()
    return render(request, 'transport/manager_drivers.html', {
        'drivers': Driver.objects.all(), 'form': form
    })


@manager_required
def manager_driver_delete(request, pk):
    get_object_or_404(Driver, pk=pk).delete()
    messages.success(request, 'تم الحذف')
    return redirect('manager_drivers')


# ── Manager: Supervisors ──────────────────────────────────────────────────────

@manager_required
def manager_supervisors(request):
    if request.method == 'POST':
        form = LineSupervisorForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'تم إضافة المسؤولة ✓')
            return redirect('manager_supervisors')
    else:
        form = LineSupervisorForm()
    return render(request, 'transport/manager_supervisors.html', {
        'supervisors': LineSupervisor.objects.all(), 'form': form
    })


@manager_required
def manager_supervisor_delete(request, pk):
    get_object_or_404(LineSupervisor, pk=pk).delete()
    messages.success(request, 'تم الحذف')
    return redirect('manager_supervisors')


# ── Manager: Employees ────────────────────────────────────────────────────────

@manager_required
def manager_employees(request):
    form = EmployeeForm()
    if request.method == 'POST':
        form = EmployeeForm(request.POST)
        if form.is_valid():
            User.objects.create_user(
                username=form.cleaned_data['username'],
                password=form.cleaned_data['password'],
                is_staff=False
            )
            messages.success(request, f'تم إنشاء حساب {form.cleaned_data["username"]} ✓')
            return redirect('manager_employees')
    employees = User.objects.filter(is_staff=False).order_by('username')
    return render(request, 'transport/manager_employees.html', {
        'employees': employees, 'form': form
    })


@manager_required
def manager_employee_delete(request, pk):
    user = get_object_or_404(User, pk=pk, is_staff=False)
    if request.method == 'POST':
        user.delete()
        messages.success(request, 'تم حذف الموظف')
        return redirect('manager_employees')
    return render(request, 'transport/confirm_delete.html', {'obj': user, 'type': 'موظف'})


# ── Manager: Excel Import ─────────────────────────────────────────────────────

@manager_required
def manager_import(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            wb = openpyxl.load_workbook(request.FILES['excel_file'])
            ws = wb.active
            count = 0
            errors = []
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]:
                    continue
                try:
                    driver = None
                    if len(row) > 5 and row[5]:
                        driver, _ = Driver.objects.get_or_create(
                            name=str(row[5]).strip(),
                            defaults={
                                'phone': str(row[6]).strip() if len(row) > 6 and row[6] else '',
                                'line_number': str(row[7]).strip() if len(row) > 7 and row[7] else '',
                            }
                        )
                    supervisor = None
                    if len(row) > 8 and row[8]:
                        supervisor, _ = LineSupervisor.objects.get_or_create(
                            name=str(row[8]).strip(),
                            defaults={
                                'phone': str(row[9]).strip() if len(row) > 9 and row[9] else '',
                                'line_number': str(row[10]).strip() if len(row) > 10 and row[10] else '',
                            }
                        )
                    student = Student.objects.create(
                        name=str(row[0]).strip(),
                        grade=str(row[1]).strip() if len(row) > 1 and row[1] else '',
                        area=str(row[2]).strip() if len(row) > 2 and row[2] else '',
                        phone1=str(row[3]).strip() if len(row) > 3 and row[3] else '',
                        phone2=str(row[4]).strip() if len(row) > 4 and row[4] else '',
                        driver=driver,
                        supervisor=supervisor,
                    )
                    for i in range(6):
                        ni = 11 + i * 2
                        pi = 12 + i * 2
                        cname = str(row[ni]).strip() if len(row) > ni and row[ni] else ''
                        cphone = str(row[pi]).strip() if len(row) > pi and row[pi] else ''
                        if cname:
                            FamilyContact.objects.create(student=student, name=cname, phone=cphone, order=i + 1)
                    count += 1
                except Exception as e:
                    errors.append(f'صف {row_num}: {e}')
            msg = f'تم استيراد {count} تلميذ بنجاح'
            if errors:
                msg += f' (أخطاء: {len(errors)})'
            messages.success(request, msg)
            for err in errors[:5]:
                messages.warning(request, err)
        except Exception as e:
            messages.error(request, f'خطأ في قراءة الملف: {e}')
        return redirect('manager_import')
    return render(request, 'transport/manager_import.html')


@manager_required
def download_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'التلاميذ'
    headers = [
        'اسم التلميذ*', 'المرحلة*', 'منطقة السكن', 'نقطة دالة', 'هاتف 1', 'هاتف 2',
        'اسم السائق', 'هاتف السائق', 'خط السائق',
        'اسم المسؤولة', 'هاتف المسؤولة', 'خط المسؤولة',
        'مخول1 اسم', 'مخول1 هاتف',
        'مخول2 اسم', 'مخول2 هاتف',
        'مخول3 اسم', 'مخول3 هاتف',
        'مخول4 اسم', 'مخول4 هاتف',
        'مخول5 اسم', 'مخول5 هاتف',
        'مخول6 اسم', 'مخول6 هاتف',
    ]
    from openpyxl.styles import Font, PatternFill, Alignment
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B2A4A', end_color='1B2A4A', fill_type='solid')
    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 18
    # Sample row
    ws.append([
        'أحمد محمد علي', 'الثالث الابتدائي', 'حي النور', 'مقابل جامع الرحمن', '07501234567', '07509876543',
        'خالد سائق', '07801111111', '5',
        'فاطمة المسؤولة', '07802222222', '5',
        'علي محمد (أب)', '07811111111',
        'هدى حسن (أم)', '07822222222',
        '', '', '', '', '', '', '', ''
    ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="student_import_template.xlsx"'
    return response


# ── Employee Search ───────────────────────────────────────────────────────────

@login_required
def employee_search(request):
    from django.db.models import Q as Qm
    lines = sorted(set(
        list(Driver.objects.exclude(line_number='').values_list('line_number', flat=True)) +
        list(LineSupervisor.objects.exclude(line_number='').values_list('line_number', flat=True))
    ))
    can_add = request.user.is_staff or request.user.groups.filter(name='can_add_students').exists()
    return render(request, 'transport/employee_search.html', {'lines': lines, 'can_add': can_add})


@login_required
def api_search(request):
    q = request.GET.get('q', '').strip()
    line = request.GET.get('line', '').strip()
    if not q and not line:
        return JsonResponse({'results': []})
    students = Student.objects.select_related('driver', 'supervisor').all()
    if q:
        students = students.filter(name__icontains=q)
    if line:
        students = students.filter(Q(driver__line_number=line) | Q(supervisor__line_number=line))
    students = students.order_by('name')[:20]
    return JsonResponse({
        'results': [{'id': s.id, 'name': s.name, 'grade': s.grade} for s in students]
    })


@login_required
def api_student(request, pk):
    s = get_object_or_404(
        Student.objects.select_related('driver', 'supervisor').prefetch_related('contacts'),
        pk=pk
    )
    contacts = list(s.contacts.order_by('order').values('name', 'phone', 'order'))
    while len(contacts) < 6:
        contacts.append({'name': '', 'phone': '', 'order': len(contacts) + 1})
    photo_url = s.photo.url if s.photo else ''
    return JsonResponse({
        'id': s.id,
        'name': s.name,
        'grade': s.grade,
        'area': s.area,
        'landmark': s.landmark,
        'phone1': s.phone1,
        'phone2': s.phone2,
        'notes': s.notes,
        'photo': photo_url,
        'driver_name': s.driver.name if s.driver else '—',
        'driver_phone': s.driver.phone if s.driver else '—',
        'driver_line': s.driver.line_number if s.driver else '—',
        'supervisor_name': s.supervisor.name if s.supervisor else '—',
        'supervisor_phone': s.supervisor.phone if s.supervisor else '—',
        'supervisor_line': s.supervisor.line_number if s.supervisor else '—',
        'contacts': contacts,
    })


# ── Single-Page Manager ───────────────────────────────────────────────────────

@manager_required
def manager_page(request):
    emp_error = None

    if request.method == 'POST':
        action = request.POST.get('action', '')
        tab = 'students'

        if action == 'add_student':
            tab = 'students'
            name = request.POST.get('name', '').strip()
            grade = request.POST.get('grade', '').strip()
            if name and grade:
                driver_id = request.POST.get('driver') or None
                supervisor_id = request.POST.get('supervisor') or None
                s = Student.objects.create(
                    name=name, grade=grade,
                    area=request.POST.get('area', '').strip(),
                    landmark=request.POST.get('landmark', '').strip(),
                    phone1=request.POST.get('phone1', '').strip(),
                    phone2=request.POST.get('phone2', '').strip(),
                    notes=request.POST.get('notes', '').strip(),
                    driver_id=driver_id, supervisor_id=supervisor_id,
                )
                if request.FILES.get('photo'):
                    s.photo = request.FILES['photo']; s.save()
                for i in range(1, 7):
                    cn = request.POST.get(f'cn{i}', '').strip()
                    cp = request.POST.get(f'cp{i}', '').strip()
                    if cn:
                        FamilyContact.objects.create(student=s, name=cn, phone=cp, order=i)
                messages.success(request, f'تم إضافة {name} ✓')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'edit_student':
            tab = 'students'
            pk = request.POST.get('student_id')
            s = get_object_or_404(Student, pk=pk)
            s.name = request.POST.get('name', s.name).strip()
            s.grade = request.POST.get('grade', s.grade).strip()
            s.area = request.POST.get('area', '').strip()
            s.landmark = request.POST.get('landmark', '').strip()
            s.phone1 = request.POST.get('phone1', '').strip()
            s.phone2 = request.POST.get('phone2', '').strip()
            s.notes = request.POST.get('notes', '').strip()
            s.driver_id = request.POST.get('driver') or None
            s.supervisor_id = request.POST.get('supervisor') or None
            if request.FILES.get('photo'):
                s.photo = request.FILES['photo']
            s.save()
            s.contacts.all().delete()
            for i in range(1, 7):
                cn = request.POST.get(f'cn{i}', '').strip()
                cp = request.POST.get(f'cp{i}', '').strip()
                if cn:
                    FamilyContact.objects.create(student=s, name=cn, phone=cp, order=i)
            messages.success(request, f'تم تحديث {s.name} ✓')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'delete_student':
            tab = 'students'
            s = get_object_or_404(Student, pk=request.POST.get('student_id'))
            messages.success(request, f'تم حذف {s.name}')
            s.delete()
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'bulk_delete_students':
            tab = 'students'
            ids = request.POST.getlist('student_ids')
            count = Student.objects.filter(id__in=ids).count()
            Student.objects.filter(id__in=ids).delete()
            messages.success(request, f'تم حذف {count} تلميذ')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'add_driver':
            tab = 'drivers'
            name = request.POST.get('drv_name', '').strip()
            phone = request.POST.get('drv_phone', '').strip()
            if name and phone:
                Driver.objects.create(name=name, phone=phone, line_number=request.POST.get('drv_line', '').strip())
                messages.success(request, f'تم إضافة السائق {name} ✓')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'delete_driver':
            tab = 'drivers'
            get_object_or_404(Driver, pk=request.POST.get('driver_id')).delete()
            messages.success(request, 'تم حذف السائق')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'add_supervisor':
            tab = 'supervisors'
            name = request.POST.get('sup_name', '').strip()
            phone = request.POST.get('sup_phone', '').strip()
            if name and phone:
                LineSupervisor.objects.create(name=name, phone=phone, line_number=request.POST.get('sup_line', '').strip())
                messages.success(request, f'تم إضافة المسؤولة {name} ✓')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'delete_supervisor':
            tab = 'supervisors'
            get_object_or_404(LineSupervisor, pk=request.POST.get('supervisor_id')).delete()
            messages.success(request, 'تم حذف المسؤولة')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'add_employee':
            tab = 'employees'
            uname = request.POST.get('emp_username', '').strip()
            pw = request.POST.get('emp_password', '')
            pw2 = request.POST.get('emp_password2', '')
            if pw != pw2:
                emp_error = 'كلمتا المرور غير متطابقتين'
            elif User.objects.filter(username=uname).exists():
                emp_error = 'اسم المستخدم موجود مسبقاً'
            elif uname and pw:
                User.objects.create_user(username=uname, password=pw, is_staff=False)
                messages.success(request, f'تم إنشاء حساب {uname} ✓')
                return redirect(f'{request.path}?tab={tab}')

        elif action == 'delete_employee':
            tab = 'employees'
            u = get_object_or_404(User, pk=request.POST.get('emp_id'), is_staff=False)
            u.delete()
            messages.success(request, 'تم حذف الموظف')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'toggle_employee':
            tab = 'employees'
            u = get_object_or_404(User, pk=request.POST.get('emp_id'), is_staff=False)
            u.is_active = not u.is_active
            u.save()
            status = 'تم تفعيل' if u.is_active else 'تم تعطيل'
            messages.success(request, f'{status} حساب {u.username}')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'reset_employee_password':
            tab = 'employees'
            u = get_object_or_404(User, pk=request.POST.get('emp_id'), is_staff=False)
            new_pw = request.POST.get('new_pw', '').strip()
            if len(new_pw) >= 4:
                u.set_password(new_pw)
                u.save()
                messages.success(request, f'تم تعديل كلمة مرور {u.username}')
            else:
                messages.error(request, 'كلمة المرور يجب أن تكون 4 أحرف على الأقل')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'toggle_employee_permission':
            tab = 'employees'
            u = get_object_or_404(User, pk=request.POST.get('emp_id'), is_staff=False)
            perm = request.POST.get('perm', '')
            from django.contrib.auth.models import Permission
            if perm == 'can_add':
                # Use groups to mark "can add" employees
                from django.contrib.auth.models import Group
                group, _ = Group.objects.get_or_create(name='can_add_students')
                if u.groups.filter(name='can_add_students').exists():
                    u.groups.remove(group)
                    messages.success(request, f'تم إلغاء صلاحية الإضافة عن {u.username}')
                else:
                    u.groups.add(group)
                    messages.success(request, f'تم منح صلاحية الإضافة لـ {u.username}')
            return redirect(f'{request.path}?tab={tab}')

        elif action == 'import_excel':
            tab = 'import'
            if request.FILES.get('excel_file'):
                try:
                    wb = openpyxl.load_workbook(request.FILES['excel_file'])
                    ws = wb.active
                    count = 0
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if not row[0]: continue
                        try:
                            driver = None
                            if len(row) > 6 and row[6]:
                                driver, _ = Driver.objects.get_or_create(name=str(row[6]).strip(),
                                    defaults={'phone': str(row[7] or '').strip(), 'line_number': str(row[8] or '').strip()})
                            supervisor = None
                            if len(row) > 9 and row[9]:
                                supervisor, _ = LineSupervisor.objects.get_or_create(name=str(row[9]).strip(),
                                    defaults={'phone': str(row[10] or '').strip(), 'line_number': str(row[11] or '').strip()})
                            st = Student.objects.create(
                                name=str(row[0]).strip(), grade=str(row[1] or '').strip(),
                                area=str(row[2] or '').strip(),
                                landmark=str(row[3] or '').strip(),
                                phone1=str(row[4] or '').strip(),
                                phone2=str(row[5] or '').strip(),
                                driver=driver, supervisor=supervisor)
                            for i in range(6):
                                ni, pi = 12 + i*2, 13 + i*2
                                cn = str(row[ni]).strip() if len(row) > ni and row[ni] else ''
                                cp = str(row[pi]).strip() if len(row) > pi and row[pi] else ''
                                if cn: FamilyContact.objects.create(student=st, name=cn, phone=cp, order=i+1)
                            count += 1
                        except Exception:
                            pass
                    messages.success(request, f'تم استيراد {count} تلميذ بنجاح ✓')
                except Exception as e:
                    messages.error(request, f'خطأ في الملف: {e}')
            return redirect(f'{request.path}?tab={tab}')

    lines = sorted(set(
        list(Driver.objects.exclude(line_number='').values_list('line_number', flat=True)) +
        list(LineSupervisor.objects.exclude(line_number='').values_list('line_number', flat=True))
    ))

    # ── Analytics data ────────────────────────────────────────────
    from django.db.models import Count
    from collections import Counter
    import json as _json

    # Students per grade
    grade_data = list(Student.objects.values('grade').annotate(count=Count('id')).order_by('-count')[:8])
    # Students per line (from driver.line_number)
    line_counts = Counter()
    for s in Student.objects.select_related('driver'):
        if s.driver and s.driver.line_number:
            line_counts[f'خط {s.driver.line_number}'] += 1
        else:
            line_counts['بدون خط'] += 1
    line_data = [{'line': k, 'count': v} for k, v in line_counts.most_common(10)]

    # Students per area (top 8)
    area_counts = Counter(s.area.strip() for s in Student.objects.exclude(area='') if s.area.strip())
    area_data = [{'area': k, 'count': v} for k, v in area_counts.most_common(8)]

    # Driver workload (students per driver)
    driver_data = [
        {'name': d.name, 'count': d.students.count()}
        for d in Driver.objects.all()[:10]
    ]
    driver_data.sort(key=lambda x: -x['count'])

    # Missing info stats
    no_driver = Student.objects.filter(driver__isnull=True).count()
    no_supervisor = Student.objects.filter(supervisor__isnull=True).count()
    no_phone = Student.objects.filter(phone1='').count()
    no_contacts = sum(1 for s in Student.objects.all() if s.contacts.count() == 0)

    context = {
        'students': Student.objects.select_related('driver', 'supervisor').prefetch_related('contacts').order_by('name'),
        'drivers': Driver.objects.all(),
        'supervisors': LineSupervisor.objects.all(),
        'employees': User.objects.filter(is_staff=False).prefetch_related('groups').order_by('username'),
        'total_students': Student.objects.count(),
        'total_drivers': Driver.objects.count(),
        'total_supervisors': LineSupervisor.objects.count(),
        'total_employees': User.objects.filter(is_staff=False, is_active=True).count(),
        'emp_error': emp_error,
        'lines': lines,
        'grade_data_json': _json.dumps(grade_data, ensure_ascii=False),
        'line_data_json': _json.dumps(line_data, ensure_ascii=False),
        'area_data_json': _json.dumps(area_data, ensure_ascii=False),
        'driver_data_json': _json.dumps(driver_data, ensure_ascii=False),
        'no_driver': no_driver,
        'no_supervisor': no_supervisor,
        'no_phone': no_phone,
        'no_contacts': no_contacts,
    }
    return render(request, 'transport/manager_page.html', context)


@login_required
def api_student_raw(request, pk):
    """Returns driver_id and supervisor_id for edit modal."""
    s = get_object_or_404(Student, pk=pk)
    return JsonResponse({'driver_id': s.driver_id, 'supervisor_id': s.supervisor_id})


# ── Error handlers ────────────────────────────────────────────────────────────

def error_404(request, exception):
    return render(request, '404.html', status=404)

def error_500(request):
    return render(request, '500.html', status=500)

def error_403(request, exception):
    return render(request, '403.html', status=403)


# ── Change Password ───────────────────────────────────────────────────────────

@manager_required
def change_password(request):
    from django.contrib.auth import update_session_auth_hash
    error = success = None
    if request.method == 'POST':
        old = request.POST.get('old_password', '')
        new1 = request.POST.get('new_password', '')
        new2 = request.POST.get('new_password2', '')
        if not request.user.check_password(old):
            error = 'كلمة المرور الحالية غير صحيحة'
        elif new1 != new2:
            error = 'كلمتا المرور الجديدتان غير متطابقتين'
        elif len(new1) < 6:
            error = 'كلمة المرور يجب أن تكون 6 أحرف على الأقل'
        else:
            request.user.set_password(new1)
            request.user.save()
            update_session_auth_hash(request, request.user)
            success = True
    return render(request, 'transport/change_password.html', {'error': error, 'success': success})


# ── Backup Database ───────────────────────────────────────────────────────────

@manager_required
def backup_database(request):
    from django.conf import settings as djsettings
    from datetime import datetime
    db_path = djsettings.DATABASES['default']['NAME']
    filename = f'backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.sqlite3'
    with open(db_path, 'rb') as f:
        response = HttpResponse(f.read(), content_type='application/octet-stream')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ── Employee with Add Permission ──────────────────────────────────────────────

@login_required
def employee_add_student(request):
    """Employees can access this ONLY if granted 'can_add_students' permission by manager."""
    if not request.user.is_staff and not request.user.groups.filter(name='can_add_students').exists():
        messages.error(request, 'ليس لديك صلاحية إضافة التلاميذ')
        return redirect('employee_search')

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        grade = request.POST.get('grade', '').strip()
        if name and grade:
            s = Student.objects.create(
                name=name, grade=grade,
                area=request.POST.get('area', '').strip(),
                landmark=request.POST.get('landmark', '').strip(),
                phone1=request.POST.get('phone1', '').strip(),
                phone2=request.POST.get('phone2', '').strip(),
                notes=request.POST.get('notes', '').strip(),
                driver_id=request.POST.get('driver') or None,
                supervisor_id=request.POST.get('supervisor') or None,
            )
            if request.FILES.get('photo'):
                s.photo = request.FILES['photo']
                s.save()
            for i in range(1, 7):
                cn = request.POST.get(f'cn{i}', '').strip()
                cp = request.POST.get(f'cp{i}', '').strip()
                if cn:
                    FamilyContact.objects.create(student=s, name=cn, phone=cp, order=i)
            messages.success(request, f'تم إضافة {name} ✓')
            return redirect('employee_search')
        else:
            messages.error(request, 'الاسم والمرحلة مطلوبان')

    return render(request, 'transport/employee_add_student.html', {
        'drivers': Driver.objects.all(),
        'supervisors': LineSupervisor.objects.all(),
    })
